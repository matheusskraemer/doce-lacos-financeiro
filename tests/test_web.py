"""Aplicação web: autenticação, telas, exportação, persistência e backup (cenários 12 a 15)."""
from __future__ import annotations

import io
import json
from datetime import timedelta
from decimal import Decimal as D

from openpyxl import load_workbook
from sqlalchemy import create_engine, select, text

from app import create_app
from app.extensions import db
from app.models import Account, LoginAttempt, Sale, Title, User
from app.services import backup
from app.services import metrics as m
from app.utils import parse_money, today

from .conftest import PASSWORD

T = today()


# 14 --------------------------------------------------------------------------
def test_rotas_protegidas_exigem_login(app, user):
    c = app.test_client()
    for url in ["/", "/vendas", "/despesas", "/contas-a-pagar", "/contas-a-receber", "/fluxo-de-caixa",
                "/relatorios", "/relatorios/vendas/exportar.csv", "/configuracoes/", "/produtos", "/compras"]:
        r = c.get(url)
        assert r.status_code == 302 and "/entrar" in r.headers["Location"], url
    r = c.post("/vendas/nova", data={"description": "x"})
    assert r.status_code == 302 and "/entrar" in r.headers["Location"]
    assert db.session.query(Sale).count() == 0


def test_csrf_obrigatorio():
    app = create_app("testing", WTF_CSRF_ENABLED=True)
    c = app.test_client()
    r = c.post("/entrar", data={"username": "x", "password": "y"})
    assert r.status_code == 400


def test_login_e_bloqueio(app, user):
    c = app.test_client()
    r = c.post("/entrar", data={"username": "carol", "password": "errada"})
    assert r.status_code == 401
    assert "Usuário ou senha inválidos" in r.get_data(as_text=True)
    for _ in range(5):
        c.post("/entrar", data={"username": "carol", "password": "errada"})
    r = c.post("/entrar", data={"username": "carol", "password": PASSWORD})
    assert r.status_code == 429  # IP bloqueado temporariamente, mesmo com a senha certa
    db.session.execute(LoginAttempt.__table__.delete())
    db.session.commit()
    r = c.post("/entrar", data={"username": "carol", "password": PASSWORD})
    assert r.status_code == 302


def test_senha_nunca_em_texto(app, user):
    u = db.session.scalar(select(User))
    assert PASSWORD not in u.password_hash and u.password_hash.startswith("scrypt:")


def test_sessao_invalidada_ao_trocar_senha(app, client):
    u = db.session.scalar(select(User))
    u.set_password("OutraSenha2026")
    db.session.commit()
    from flask import g
    g.pop("_login_user", None)  # nos testes o contexto da aplicação é compartilhado entre requisições
    r = client.get("/")
    assert r.status_code == 302


def test_cabecalhos_de_seguranca(client):
    r = client.get("/")
    assert "default-src 'self'" in r.headers["Content-Security-Policy"]
    assert r.headers["X-Frame-Options"] == "DENY"
    assert r.headers["Cache-Control"] == "no-store"


def test_redirect_seguro_apos_login(app, user):
    c = app.test_client()
    r = c.post("/entrar?next=https://malicioso.example/x", data={"username": "carol", "password": PASSWORD})
    assert r.headers["Location"] == "/"


def test_todas_as_telas_abrem_vazias(client):
    for url in ["/", "/vendas", "/vendas/nova", "/receitas", "/receitas/nova", "/despesas", "/despesas/nova",
                "/produtos", "/produtos/novo", "/compras", "/compras/nova", "/contas-a-pagar", "/contas-a-receber",
                "/fluxo-de-caixa", "/fluxo-de-caixa?visao=semana", "/fluxo-de-caixa?visao=mes", "/transferencias",
                "/relatorios", "/configuracoes/", "/marca.css", "/marca/logo"]:
        r = client.get(url)
        assert r.status_code == 200, url
    for tab in ["empresa", "contas", "categorias", "formas", "seguranca", "backup", "historico"]:
        assert client.get(f"/configuracoes/?aba={tab}").status_code == 200
    html = client.get("/").get_data(as_text=True)
    assert "Comece cadastrando suas contas" in html
    assert "R$ " in html or "R$ 0,00" in html


def _seed_via_web(client):
    client.post("/configuracoes/contas", data={"name": "Banco", "kind": "banco", "opening_balance": "1.000,00",
                                               "opening_date": (T - timedelta(days=30)).isoformat(), "include_in_total": "1"})
    acc = db.session.scalar(select(Account.id).where(Account.name == "Banco"))
    r = client.post("/vendas/nova", data={
        "description": "Kit escolar", "sale_date": T.isoformat(), "customer": "Escola X",
        "item_product": ["", ""], "item_desc": ["Laço azul", "Tiara"], "item_qty": ["10", "5"],
        "item_price": ["12,50", "20,00"], "discount": "25,00", "received_now": "1", "account_id": str(acc),
        "request_key": "k1"})
    assert r.status_code == 302
    r = client.post("/despesas/nova", data={
        "description": "Frete Correios", "category_id": "8", "amount": "35,90", "competence_date": T.isoformat(),
        "due_date": (T + timedelta(days=3)).isoformat(), "settle_now": "", "installments": "1", "recurrence": "1",
        "request_key": "k2"})
    assert r.status_code == 302
    return acc


def test_fluxo_completo_pela_interface(client):
    acc = _seed_via_web(client)
    sale = db.session.scalar(select(Sale))
    assert sale.gross_amount == D("225.00") and sale.net_amount == D("200.00")
    assert m.total_balance(T) == D("1200.00")
    # reenvio do mesmo formulário não duplica
    client.post("/vendas/nova", data={"description": "Kit escolar", "sale_date": T.isoformat(), "gross_amount": "200",
                                      "received_now": "1", "account_id": str(acc), "request_key": "k1"})
    assert db.session.query(Sale).count() == 1
    t = db.session.scalar(select(Title).where(Title.kind == "pagar"))
    r = client.post(f"/lancamentos/{t.id}/liquidar", data={"amount": "35,90", "date": T.isoformat(),
                                                           "account_id": str(acc), "request_key": "k3"})
    assert r.status_code == 302
    assert m.total_balance(T) == D("1164.10")
    for url in ["/", "/vendas", f"/vendas/{sale.id}", f"/lancamentos/{t.id}", "/contas-a-pagar?ver=todas",
                "/fluxo-de-caixa", "/despesas", "/configuracoes/?aba=historico"]:
        assert client.get(url).status_code == 200, url
    html = client.get("/").get_data(as_text=True)
    assert "1.164,10" in html


# 12 --------------------------------------------------------------------------
def test_exportacao_respeita_filtros(client):
    _seed_via_web(client)
    from app.services.reports import CATALOG
    for key, _, _ in CATALOG:
        assert client.get(f"/relatorios/{key}?periodo=mes").status_code == 200, key
        r = client.get(f"/relatorios/{key}/exportar.csv?periodo=mes")
        assert r.status_code == 200 and r.data.startswith("﻿".encode()), key
        x = client.get(f"/relatorios/{key}/exportar.xlsx?periodo=mes")
        assert x.status_code == 200
        load_workbook(io.BytesIO(x.data))
    csv = client.get("/relatorios/vendas/exportar.csv?periodo=mes").data.decode("utf-8-sig")
    assert "Kit escolar" in csv and "200,00" in csv and T.strftime("%d/%m/%Y") in csv
    # filtro que exclui tudo -> sem a venda
    old = T - timedelta(days=400)
    csv2 = client.get(f"/relatorios/vendas/exportar.csv?periodo=personalizado&inicio={old}&fim={old}").data.decode("utf-8-sig")
    assert "Kit escolar" not in csv2
    wb = load_workbook(io.BytesIO(client.get("/relatorios/vendas/exportar.xlsx?periodo=mes").data))
    ws = wb.active
    values = [c.value for row in ws.iter_rows() for c in row]
    assert 200.0 in values


# 13 --------------------------------------------------------------------------
def test_dados_persistem_apos_reabrir(app, client):
    _seed_via_web(client)
    db.session.remove()
    db.engine.dispose()
    app2 = create_app("testing")
    with app2.app_context():
        assert db.session.query(Sale).count() == 1
        assert m.total_balance(T) == D("1200.00")


# 15 --------------------------------------------------------------------------
def test_backup_e_restauracao(app, client):
    _seed_via_web(client)
    payload = json.loads(backup.dumps(backup.export_all()))
    expected_balance = m.total_balance(T)
    expected_open = m.open_total("pagar")

    # restaura em outro banco (vazio) e compara
    url = app.config["SQLALCHEMY_DATABASE_URI"].rsplit("/", 1)[0] + "/docelacos_restore"
    app_r = create_app("testing", SQLALCHEMY_DATABASE_URI=url)
    with app_r.app_context():
        db.drop_all()
        db.create_all()
        counts = backup.restore_all(payload)
        assert counts["sales"] == 1
        assert m.total_balance(T) == expected_balance
        assert m.open_total("pagar") == expected_open
        assert m.integrity_report() == []
        # sequências ajustadas: novos registros não colidem
        db.session.add(Account(name="Nova", opening_balance=0, opening_date=T))
        db.session.commit()
        # login funciona com a senha original
        c = app_r.test_client()
        assert c.post("/entrar", data={"username": "carol", "password": PASSWORD}).status_code == 302
        # recusa sobrescrever sem --force e recusa arquivo adulterado
        import pytest
        with pytest.raises(backup.RestoreError):
            backup.restore_all(payload)
        bad = json.loads(json.dumps(payload))
        bad["tables"]["sales"][0]["net_amount"] = "1.00"
        with pytest.raises(backup.RestoreError):
            backup.restore_all(bad, force=True)
        db.session.remove()
        db.drop_all()


def test_download_backup_exige_senha(client):
    r = client.post("/configuracoes/backup", data={"password": "errada"})
    assert r.status_code == 302
    r = client.post("/configuracoes/backup", data={"password": PASSWORD})
    assert r.status_code == 200 and r.mimetype == "application/json"
    assert json.loads(r.data)["format"] == "doce-lacos-backup"


def test_parse_money_formatos_brasileiros():
    assert parse_money("1.234,56") == D("1234.56")
    assert parse_money("R$ 10") == D("10.00")
    assert parse_money("12.5") == D("12.50")
    assert parse_money("1.234") == D("1234.00")
    assert parse_money("0,1") == D("0.10")
    import pytest
    from app.utils import ValidationError
    for bad in ["abc", "1,2,3", "-5", "0", "1,234"]:
        if bad == "1,234":
            with pytest.raises(ValidationError):
                parse_money(bad)
            continue
        with pytest.raises(ValidationError):
            parse_money(bad)


def test_logo_svg_com_script_e_recusado(client):
    data = {"company_name": "Doce Laços", "color_primary": "#8E4A67", "color_accent": "#B8915A",
            "color_ink": "#2E2433", "projection_days": "30", "upcoming_days": "7",
            "logo": (io.BytesIO(b"<svg><script>alert(1)</script></svg>"), "logo.svg", "image/svg+xml")}
    client.post("/configuracoes/empresa", data=data, content_type="multipart/form-data")
    from app.models import CompanySettings
    assert db.session.get(CompanySettings, 1).logo_data is None


def test_ensure_user_por_variaveis_de_ambiente(app, monkeypatch):
    runner = app.test_cli_runner()
    r = runner.invoke(args=["ensure-user"])
    assert "Defina INITIAL_ADMIN_USERNAME" in r.output
    monkeypatch.setenv("INITIAL_ADMIN_USERNAME", "Carolina")
    monkeypatch.setenv("INITIAL_ADMIN_PASSWORD", "fraca")
    r = runner.invoke(args=["ensure-user"])
    assert r.exit_code == 1 and db.session.query(User).count() == 0
    monkeypatch.setenv("INITIAL_ADMIN_PASSWORD", PASSWORD)
    r = runner.invoke(args=["ensure-user"])
    assert r.exit_code == 0 and PASSWORD not in r.output
    u = db.session.scalar(select(User))
    assert u.username == "carolina" and u.check_password(PASSWORD)
    r = runner.invoke(args=["ensure-user"])  # não cria outro nem altera a senha
    assert db.session.query(User).count() == 1 and "Remova INITIAL_ADMIN_PASSWORD" in r.output
