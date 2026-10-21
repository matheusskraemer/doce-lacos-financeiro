"""Relatórios: cada relatório devolve colunas, linhas e totais a partir dos filtros.

As mesmas funções alimentam a tela e as exportações (CSV/Excel), garantindo que o
arquivo exportado respeite exatamente os filtros aplicados.
"""
from __future__ import annotations

import csv
import io
from dataclasses import dataclass, field
from datetime import date
from decimal import Decimal

from sqlalchemy import or_, select

from ..extensions import db
from ..models import Category, Purchase, Sale, Settlement, Title, Transfer
from ..utils import ZERO, add_months, fmt_date, month_label, month_start, q, today
from . import metrics as m


@dataclass
class Column:
    key: str
    label: str
    kind: str = "text"  # text | money | date | number


@dataclass
class Report:
    key: str
    title: str
    description: str
    columns: list[Column]
    rows: list[dict] = field(default_factory=list)
    totals: dict = field(default_factory=dict)
    notes: list[str] = field(default_factory=list)
    summary: list[tuple[str, Decimal]] = field(default_factory=list)


@dataclass
class Filters:
    start: date
    end: date
    category_id: int | None = None
    account_id: int | None = None
    status: str = ""
    search: str = ""
    basis: str = "caixa"
    month_a: date | None = None
    month_b: date | None = None


CATALOG = [
    ("vendas", "Vendas por período", "Vendas registradas, com recebido e a receber."),
    ("recebimentos", "Receitas recebidas", "Tudo o que efetivamente entrou nas contas."),
    ("pagamentos", "Despesas pagas", "Tudo o que efetivamente saiu das contas."),
    ("despesas_categoria", "Despesas por categoria", "Totais por categoria, por caixa ou competência."),
    ("compras", "Compras de mercadorias", "Compras para revenda e situação de pagamento."),
    ("a_pagar", "Contas a pagar", "Compromissos em aberto e já liquidados."),
    ("a_receber", "Contas a receber", "Valores a receber em aberto e já recebidos."),
    ("fluxo_mensal", "Fluxo de caixa mensal", "Entradas, saídas e saldo ao fim de cada mês."),
    ("comparativo", "Comparativo entre meses", "Dois meses lado a lado."),
    ("movimentacoes", "Histórico de movimentações", "Liquidações, estornos e transferências."),
    ("resultado", "Resultado financeiro", "Resultado de caixa e visão gerencial por competência."),
]
CATALOG_KEYS = {k for k, _, _ in CATALOG}


def _settlement_rows(kind: str, flt: Filters) -> Report:
    stmt = (select(Settlement).join(Title, Title.id == Settlement.title_id)
            .where(Title.kind == kind, Settlement.date.between(flt.start, flt.end))
            .order_by(Settlement.date, Settlement.id))
    if flt.category_id:
        stmt = stmt.where(Title.category_id == flt.category_id)
    if flt.account_id:
        stmt = stmt.where(Settlement.account_id == flt.account_id)
    if flt.search:
        like = f"%{flt.search}%"
        stmt = stmt.where(or_(Title.description.ilike(like), Title.counterparty.ilike(like)))
    rows = []
    total = ZERO
    for s in db.session.scalars(stmt):
        val = s.signed_amount
        total += val
        rows.append({
            "date": s.date, "description": s.title.description, "category": s.title.category.name,
            "counterparty": s.title.counterparty or "", "account": s.account.name,
            "method": s.payment_method.name if s.payment_method else "",
            "type": "Estorno" if s.is_reversal else ("Recebimento" if kind == "receber" else "Pagamento"),
            "amount": val, "_url": ("titulos.detail", {"title_id": s.title_id}),
        })
    noun = "recebido" if kind == "receber" else "pago"
    return Report(
        key="", title="", description="",
        columns=[Column("date", "Data", "date"), Column("description", "Descrição"), Column("category", "Categoria"),
                 Column("counterparty", "Cliente" if kind == "receber" else "Fornecedor"), Column("account", "Conta"),
                 Column("method", "Forma"), Column("type", "Tipo"), Column("amount", "Valor", "money")],
        rows=rows, totals={"amount": q(total)},
        notes=[f"Valores pela data efetiva da movimentação. Estornos aparecem negativos e já estão descontados do total {noun}."],
    )


def build(key: str, flt: Filters) -> Report:
    title, desc = next(((t, d) for k, t, d in CATALOG if k == key), (None, None))
    if title is None:
        raise KeyError(key)
    ref = today()

    if key == "vendas":
        stmt = select(Sale).where(Sale.sale_date.between(flt.start, flt.end)).order_by(Sale.sale_date, Sale.id)
        if flt.search:
            like = f"%{flt.search}%"
            stmt = stmt.where(or_(Sale.description.ilike(like), Sale.customer.ilike(like)))
        rows, tot = [], {"gross": ZERO, "discount": ZERO, "net": ZERO, "received": ZERO, "to_receive": ZERO}
        labels = {"recebida": "Recebida", "parcial": "Recebida em parte", "a_receber": "A receber", "cancelada": "Cancelada"}
        for s in db.session.scalars(stmt):
            st = s.receipt_status
            if flt.status and st != flt.status:
                continue
            r = {"date": s.sale_date, "description": s.description, "customer": s.customer or "",
                 "gross": s.gross_amount, "discount": s.discount, "net": s.net_amount,
                 "received": s.received, "to_receive": s.to_receive, "status": labels[st],
                 "_url": ("vendas.detail", {"sale_id": s.id}), "_muted": s.status == "cancelada"}
            rows.append(r)
            if s.status == "ativa":
                for k in tot:
                    tot[k] += r[k]
        rep = Report(key, title, desc, [
            Column("date", "Data", "date"), Column("description", "Venda"), Column("customer", "Cliente"),
            Column("gross", "Bruto", "money"), Column("discount", "Desconto", "money"), Column("net", "Líquido", "money"),
            Column("received", "Recebido", "money"), Column("to_receive", "A receber", "money"), Column("status", "Situação")],
            rows, {k: q(v) for k, v in tot.items()},
            ["Vendas canceladas aparecem na lista mas não entram nos totais."])
        return rep

    if key in ("recebimentos", "pagamentos"):
        rep = _settlement_rows("receber" if key == "recebimentos" else "pagar", flt)
        rep.key, rep.title, rep.description = key, title, desc
        return rep

    if key == "despesas_categoria":
        basis = flt.basis if flt.basis in ("caixa", "competencia") else "caixa"
        data = m.paid_by_category(flt.start, flt.end, basis)
        total = q(sum((d["total"] for d in data), ZERO))
        rows = [{"category": d["name"], "group": "Mercadorias" if d["group"] == "mercadoria" else "Operacional",
                 "total": d["total"], "share": f"{(d['total'] / total * 100):.1f}%".replace(".", ",") if total else "—"}
                for d in data]
        return Report(key, title, desc, [Column("category", "Categoria"), Column("group", "Tipo"),
                                         Column("total", "Total", "money"), Column("share", "Participação")],
                      rows, {"total": total},
                      ["Base caixa: valores efetivamente pagos no período." if basis == "caixa"
                       else "Base competência: despesas lançadas no período, pagas ou não (exceto canceladas)."])

    if key == "compras":
        stmt = select(Purchase).where(Purchase.purchase_date.between(flt.start, flt.end)).order_by(Purchase.purchase_date)
        if flt.search:
            stmt = stmt.where(Purchase.supplier.ilike(f"%{flt.search}%"))
        rows, tot = [], {"total": ZERO, "paid": ZERO, "to_pay": ZERO}
        labels = {"paga": "Paga", "parcial": "Paga em parte", "a_pagar": "A pagar", "cancelada": "Cancelada"}
        for p in db.session.scalars(stmt):
            st = p.payment_status
            if flt.status and st != flt.status:
                continue
            items = ", ".join(f"{i.description} ×{i.quantity.normalize():f}".replace(".", ",") for i in p.items)
            r = {"date": p.purchase_date, "supplier": p.supplier or "", "items": items, "total": p.total,
                 "paid": p.paid, "to_pay": p.to_pay, "status": labels[st],
                 "_url": ("produtos.compra", {"purchase_id": p.id}), "_muted": p.status == "cancelada"}
            rows.append(r)
            if p.status == "ativa":
                for k in tot:
                    tot[k] += r[k]
        return Report(key, title, desc, [Column("date", "Data", "date"), Column("supplier", "Fornecedor"),
                                         Column("items", "Itens"), Column("total", "Total", "money"),
                                         Column("paid", "Pago", "money"), Column("to_pay", "A pagar", "money"),
                                         Column("status", "Situação")],
                      rows, {k: q(v) for k, v in tot.items()},
                      ["A compra entra como saída de caixa somente quando é paga (ver Despesas pagas)."])

    if key in ("a_pagar", "a_receber"):
        kind = "pagar" if key == "a_pagar" else "receber"
        stmt = (select(Title).where(Title.kind == kind, Title.due_date.between(flt.start, flt.end))
                .order_by(Title.due_date, Title.id))
        if flt.category_id:
            stmt = stmt.where(Title.category_id == flt.category_id)
        if flt.search:
            like = f"%{flt.search}%"
            stmt = stmt.where(or_(Title.description.ilike(like), Title.counterparty.ilike(like)))
        labels = {"vencido": "Vencida", "pendente": "A vencer", "parcial": "Em parte", "liquidado": "Liquidada",
                  "cancelado": "Cancelada"}
        rows, tot = [], {"amount": ZERO, "settled": ZERO, "open": ZERO}
        for t in db.session.scalars(stmt):
            st = t.status_on(ref)
            if flt.status and st != flt.status:
                continue
            r = {"due": t.due_date, "description": t.description + (f" ({t.installment_label})" if t.installment_label and t.origin not in ("venda", "compra") else ""),
                 "counterparty": t.counterparty or "", "category": t.category.name, "amount": t.amount,
                 "settled": t.settled_amount, "open": t.open_amount, "status": labels[st],
                 "_url": ("titulos.detail", {"title_id": t.id}), "_muted": t.is_canceled}
            rows.append(r)
            if not t.is_canceled:
                for k in tot:
                    tot[k] += q(r[k])
        return Report(key, title, desc, [Column("due", "Vencimento", "date"), Column("description", "Descrição"),
                                         Column("counterparty", "Fornecedor" if kind == "pagar" else "Cliente"),
                                         Column("category", "Categoria"), Column("amount", "Valor", "money"),
                                         Column("settled", "Liquidado", "money"), Column("open", "Em aberto", "money"),
                                         Column("status", "Situação")],
                      rows, {k: q(v) for k, v in tot.items()}, ["Filtrado pela data de vencimento."])

    if key == "fluxo_mensal":
        rows = []
        cur = month_start(flt.start)
        tot = {"received": ZERO, "paid": ZERO, "result": ZERO}
        while cur <= flt.end:
            sm = m.month_summary(cur)
            rows.append({"month": month_label(cur).capitalize(), "received": sm.received, "paid": sm.paid,
                         "operational": sm.paid_operational, "goods": sm.paid_goods, "result": sm.result,
                         "closing": sm.closing_balance if sm.closing_balance is not None else ""})
            for k in tot:
                tot[k] += getattr(sm, k)
            cur = add_months(cur, 1)
        return Report(key, title, desc, [Column("month", "Mês"), Column("received", "Entradas", "money"),
                                         Column("paid", "Saídas", "money"),
                                         Column("operational", "Saídas operacionais", "money"),
                                         Column("goods", "Pagto. mercadorias", "money"),
                                         Column("result", "Resultado de caixa", "money"),
                                         Column("closing", "Saldo ao fim do mês", "money")],
                      rows, {k: q(v) for k, v in tot.items()},
                      ["Resultado de caixa = entradas − saídas efetivadas. Não é lucro contábil."])

    if key == "comparativo":
        a = month_start(flt.month_a or add_months(ref, -1))
        b = month_start(flt.month_b or ref)
        sa, sb = m.month_summary(a), m.month_summary(b)
        lines = [("Vendas (líquidas)", "sales"), ("Recebido", "received"), ("Pago", "paid"),
                 ("Despesas operacionais pagas", "paid_operational"), ("Mercadorias pagas", "paid_goods"),
                 ("Compras de mercadorias (registradas)", "purchases"), ("Resultado de caixa", "result")]
        rows = []
        for label, attr in lines:
            va, vb = getattr(sa, attr), getattr(sb, attr)
            diff = q(vb - va)
            pct = f"{(diff / abs(va) * 100):+.1f}%".replace(".", ",") if va else "—"
            rows.append({"metric": label, "a": va, "b": vb, "diff": diff, "pct": pct})
        return Report(key, title, desc, [Column("metric", "Indicador"),
                                         Column("a", month_label(a).capitalize(), "money"),
                                         Column("b", month_label(b).capitalize(), "money"),
                                         Column("diff", "Diferença", "money"), Column("pct", "Variação")],
                      rows, {}, [])

    if key == "movimentacoes":
        rows = []
        s_stmt = (select(Settlement).join(Title, Title.id == Settlement.title_id)
                  .where(Settlement.date.between(flt.start, flt.end)))
        if flt.account_id:
            s_stmt = s_stmt.where(Settlement.account_id == flt.account_id)
        for s in db.session.scalars(s_stmt):
            kind = "Entrada" if s.title.kind == "receber" else "Saída"
            rows.append({"date": s.date, "type": ("Estorno de " + kind.lower()) if s.is_reversal else kind,
                         "description": s.title.description, "account": s.account.name, "effect": s.cash_effect,
                         "registered": s.created_at, "_url": ("titulos.detail", {"title_id": s.title_id}), "_id": s.id})
        t_stmt = select(Transfer).where(Transfer.date.between(flt.start, flt.end))
        for t in db.session.scalars(t_stmt):
            for acc, eff in ((t.from_account, -t.amount), (t.to_account, t.amount)):
                if flt.account_id and acc.id != flt.account_id:
                    continue
                rows.append({"date": t.date, "type": "Estorno de transferência" if t.reverses_id else "Transferência",
                             "description": f"{t.from_account.name} → {t.to_account.name}" + (f" · {t.notes}" if t.notes else ""),
                             "account": acc.name, "effect": eff, "registered": t.created_at,
                             "_url": ("fluxo.transferencias", {}), "_id": t.id})
        rows.sort(key=lambda r: (r["date"], r["registered"]))
        if flt.search:
            s = flt.search.lower()
            rows = [r for r in rows if s in r["description"].lower()]
        total = q(sum((r["effect"] for r in rows), ZERO))
        return Report(key, title, desc, [Column("date", "Data", "date"), Column("type", "Tipo"),
                                         Column("description", "Descrição"), Column("account", "Conta"),
                                         Column("effect", "Efeito no saldo", "money"),
                                         Column("registered", "Registrado em", "date")],
                      rows, {"effect": total},
                      ["Inclui todas as contas (também as fora do saldo consolidado), salvo filtro por conta."])

    if key == "resultado":
        cash = m.cash_totals(flt.start, flt.end)
        sales = m.sales_totals(flt.start, flt.end)
        other = m.other_revenue_competence(flt.start, flt.end)
        exp = m.expenses_competence(flt.start, flt.end)
        managerial = q(sales["net"] + other - exp["operacional"] - exp["mercadoria"])
        rows = [
            {"group": "Caixa", "line": "Recebimentos efetivados", "value": cash["in"]},
            {"group": "Caixa", "line": "(−) Pagamentos efetivados", "value": -cash["out"]},
            {"group": "Caixa", "line": "= Resultado de caixa", "value": cash["result"], "_strong": True},
            {"group": "Competência", "line": "Vendas líquidas (bruto − descontos)", "value": sales["net"]},
            {"group": "Competência", "line": "Outras receitas lançadas", "value": other},
            {"group": "Competência", "line": "(−) Despesas operacionais lançadas", "value": -exp["operacional"]},
            {"group": "Competência", "line": "(−) Compras de mercadorias lançadas", "value": -exp["mercadoria"]},
            {"group": "Competência", "line": "= Resultado gerencial simplificado", "value": managerial, "_strong": True},
        ]
        return Report(key, title, desc, [Column("group", "Visão"), Column("line", "Linha"),
                                         Column("value", "Valor", "money")], rows, {},
                      ["O resultado de caixa mostra quanto o dinheiro disponível aumentou ou diminuiu.",
                       "O resultado gerencial simplificado considera as compras de mercadorias do período inteiras, "
                       "não o custo das mercadorias vendidas, nem tributos sobre o lucro. Não é lucro contábil."])
    raise KeyError(key)


# --------------------------------------------------------------------------- exportação
def _cell_text(value, kind: str) -> str:
    if value is None or value == "":
        return ""
    if kind == "money" or isinstance(value, Decimal):
        return f"{q(value):.2f}".replace(".", ",")
    if kind == "date" or isinstance(value, date):
        return fmt_date(value)
    return str(value)


def to_csv(rep: Report) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf, delimiter=";", quoting=csv.QUOTE_MINIMAL)
    w.writerow([c.label for c in rep.columns])
    for r in rep.rows:
        w.writerow([_cell_text(r.get(c.key), c.kind) for c in rep.columns])
    if rep.totals:
        w.writerow([("Total" if i == 0 else _cell_text(rep.totals.get(c.key), c.kind)) for i, c in enumerate(rep.columns)])
    # BOM para o Excel reconhecer acentos
    return ("﻿" + buf.getvalue()).encode("utf-8")


def to_xlsx(rep: Report, subtitle: str = "") -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = rep.title[:31]
    ws.append([rep.title])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([subtitle])
    ws.append([])
    ws.append([c.label for c in rep.columns])
    head_row = ws.max_row
    for i in range(1, len(rep.columns) + 1):
        cell = ws.cell(head_row, i)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="5B3A4C")
        cell.alignment = Alignment(vertical="center")
    for r in rep.rows:
        values = []
        for c in rep.columns:
            v = r.get(c.key)
            if v == "":
                v = None
            if c.kind == "money" and v is not None:
                v = float(q(v))
            elif hasattr(v, "tzinfo") and v is not None and getattr(v, "tzinfo", None):
                v = v.replace(tzinfo=None)
            values.append(v)
        ws.append(values)
    if rep.totals:
        ws.append([("Total" if i == 0 else (float(rep.totals[c.key]) if c.key in rep.totals else None))
                   for i, c in enumerate(rep.columns)])
        for i in range(1, len(rep.columns) + 1):
            ws.cell(ws.max_row, i).font = Font(bold=True)
    for i, c in enumerate(rep.columns, start=1):
        letter = get_column_letter(i)
        for row in ws.iter_rows(min_row=head_row + 1, min_col=i, max_col=i):
            for cell in row:
                if c.kind == "money":
                    cell.number_format = '"R$" #,##0.00;[Red]-"R$" #,##0.00'
                elif c.kind == "date":
                    cell.number_format = "DD/MM/YYYY"
        width = max([len(c.label)] + [len(_cell_text(r.get(c.key), c.kind)) for r in rep.rows[:500]] + [8])
        ws.column_dimensions[letter].width = min(width + 2, 60)
    ws.freeze_panes = ws.cell(head_row + 1, 1)
    if rep.notes:
        ws.append([])
        for n in rep.notes:
            ws.append([n])
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
